"""Reading an FIR document and pulling structured records out of it.

This is rule-based information extraction, not a neural model, and it says so:
every value it produces carries the span of text it came from, so an officer
reviewing the parse can see *why* the system believes a number is an IMEI
rather than being asked to trust it. That provenance is the same contract the
detection mechanisms work under - a value with no retrievable source is a bug,
not a low-confidence result.

Three layers, in decreasing reliability:

1. **Labelled fields.** An FIR is a form. `District : Jaipur` is not a guess.
2. **Identifier patterns.** An IMEI is fifteen digits, a registration is
   `RJ 14 AB 1234`. These are shapes, and shapes can be matched exactly.
3. **A modus operandi lexicon.** `gas cutter` and `oxy-acetylene torch` are the
   same method described twice, which is the whole reason the MO mechanism
   carries a structured vector rather than working off the narrative text.

Layer 3 is the only one that infers anything, and it is the one the review step
exists for. What this module does not do is resolve names against a gazetteer
or disambiguate people - that is entity resolution, and it happens afterwards,
in `er.py`, where it is explainable.
"""
from __future__ import annotations

import io
import re
from typing import Any

# --------------------------------------------------------------------------
# Document text
# --------------------------------------------------------------------------

SUPPORTED = (".txt", ".text", ".md", ".docx", ".pdf")


class ExtractionError(ValueError):
    """A document this module cannot read."""


def read_document(filename: str, data: bytes) -> tuple[str, str]:
    """Return the document's text and a note on how it was obtained."""
    name = (filename or "").lower()
    if name.endswith((".txt", ".text", ".md")):
        for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
            try:
                return data.decode(encoding), f"plain text ({encoding})"
            except UnicodeDecodeError:
                continue
        raise ExtractionError("the text file is not in a recognised encoding")

    if name.endswith(".docx"):
        try:
            import docx
        except ImportError:
            raise ExtractionError("python-docx is not installed") from None
        try:
            doc = docx.Document(io.BytesIO(data))
        except Exception:
            raise ExtractionError("the file is not a readable .docx") from None
        parts = [p.text for p in doc.paragraphs]
        # FIR forms put the particulars in a table and the narrative below it,
        # so a reader that only walks paragraphs loses the half that matters.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" : ".join(c for c in cells if c))
        return "\n".join(parts), "Word document"

    if name.endswith(".pdf"):
        try:
            import fitz  # PyMuPDF
        except ImportError:
            fitz = None
        if fitz is not None:
            try:
                with fitz.open(stream=data, filetype="pdf") as doc:
                    text = "\n".join(page.get_text() for page in doc)
                if text.strip():
                    return text, "PDF text layer"
            except Exception:
                pass
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ExtractionError("no PDF reader is installed") from None
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(p.extract_text() or "" for p in reader.pages)
        except Exception:
            raise ExtractionError("the file is not a readable PDF") from None
        if not text.strip():
            raise ExtractionError(
                "this PDF has no text layer - it is a scan, and OCR is not in "
                "this build")
        return text, "PDF text layer"

    raise ExtractionError(
        f"unsupported file type; this build reads {', '.join(SUPPORTED)}")


# --------------------------------------------------------------------------
# Provenance
# --------------------------------------------------------------------------

def _found(value: Any, span: tuple[int, int] | None, text: str,
           how: str) -> dict[str, Any]:
    """One extracted value, with the text it came from."""
    evidence = ""
    if span:
        lo = max(0, span[0] - 42)
        hi = min(len(text), span[1] + 42)
        evidence = ("…" if lo else "") + text[lo:hi].strip().replace("\n", " ") \
                   + ("…" if hi < len(text) else "")
    return {"value": value, "how": how, "evidence": evidence,
            "span": list(span) if span else None}


# --------------------------------------------------------------------------
# Layer 1 - labelled fields
# --------------------------------------------------------------------------

def _labelled(text: str, *names: str) -> tuple[str, tuple[int, int]] | None:
    """The value of the first form field matching any of these labels."""
    for name in names:
        pattern = re.compile(
            rf"^[ \t]*{name}[ \t]*[:\-–]+[ \t]*(?P<v>.+?)[ \t]*$",
            re.IGNORECASE | re.MULTILINE)
        m = pattern.search(text)
        if m and m.group("v").strip(" .-"):
            return m.group("v").strip(" .-"), m.span("v")
    return None


def _narrative(text: str) -> tuple[str, tuple[int, int]] | None:
    """The body of the FIR.

    Preferred from a labelled section, because that is what the form says it is.
    Failing that, the longest run of prose in the document - an FIR's particulars
    are short labelled lines and its account of the offence is not.
    """
    heading = re.compile(
        r"^[ \t]*(?:brief\s+facts|facts\s+of\s+the\s+case|statement|"
        r"complaint|narrative|contents?\s+of\s+(?:the\s+)?fir)"
        r"[ \t]*[:\-–]*[ \t]*$",
        re.IGNORECASE | re.MULTILINE)
    m = heading.search(text)
    if m:
        rest = text[m.end():]
        # The account of the offence runs until the form resumes. A resumed
        # form looks like a short label followed by a colon at the start of a
        # line, which prose does not: without this the narrative swallows
        # "Property stolen :", the valuation and the station's signature block.
        stop = re.search(r"\n[ \t]*[A-Z][A-Za-z .,/'()-]{2,38}[ \t]*[:–]", rest)
        body = (rest[: stop.start()] if stop else rest).strip()
        if len(body) > 80:
            start = m.end() + rest.index(body[:20])
            return body, (start, start + len(body))

    inline = _labelled(text, r"brief\s+facts", r"facts\s+of\s+the\s+case",
                       "narrative", "complaint")
    if inline and len(inline[0]) > 80:
        return inline

    best, best_span = "", None
    for block in re.finditer(r"[^\n]{100,}(?:\n[^\n]{20,})*", text):
        chunk = block.group(0).strip()
        if len(chunk) > len(best):
            best, best_span = chunk, block.span()
    return (best, best_span) if len(best) > 80 else None


# --------------------------------------------------------------------------
# Layer 2 - identifier shapes
# --------------------------------------------------------------------------

# Order matters. An IMEI is fifteen digits and a phone number is ten, so the
# IMEI is taken first and its digits are masked out before phones are looked
# for - otherwise ten digits of an IMEI become a phone number that was never
# in the document.
IMEI_RE = re.compile(r"(?<!\d)(\d{15})(?!\d)")
PHONE_RE = re.compile(r"(?<![\d/])(?:\+?91[\s-]?)?([6-9]\d{4}[\s-]?\d{5})(?!\d)")
ACCOUNT_RE = re.compile(
    r"(?:a\s*/\s*c|account|acct)[^\n:]{0,20}[:\-–\s]+(?:no\.?\s*)?"
    r"[:\-–\s]*(\d{9,18})", re.IGNORECASE)
REG_RE = re.compile(
    r"(?<![A-Z0-9])([A-Z]{2}[\s-]?\d{1,2}[\s-]?[A-Z]{1,3}[\s-]?\d{4})(?![A-Z0-9])")
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})(?!\d)")
ISO_DATE_RE = re.compile(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)")


def _normalise_date(value: str) -> str | None:
    m = ISO_DATE_RE.search(value)
    if m:
        return m.group(0)
    m = DATE_RE.search(value)
    if not m:
        return None
    day, month, year = m.groups()
    year = int(year)
    if year < 100:
        year += 2000
    try:
        if not (1 <= int(day) <= 31 and 1 <= int(month) <= 12):
            return None
    except ValueError:
        return None
    return f"{year:04d}-{int(month):02d}-{int(day):02d}"


# --------------------------------------------------------------------------
# Layer 3 - the modus operandi lexicon
# --------------------------------------------------------------------------

# Each entry maps a phrase to one value of one structured attribute. The values
# are exactly those already recorded against other cases: a parse that invented
# a new value would produce a vector the mechanism has nothing to compare to.
MO_LEXICON: list[tuple[str, str, str]] = [
    ("tool_class", "thermal_cutting",
     r"gas\s*cutter|oxy-?\s*acetylene|cutting\s+torch|blow\s*torch|gas\s+cutting"),
    ("tool_class", "mechanical_force",
     r"crow\s*bar|jemmy|iron\s+rod|hammer|screw\s*driver|brute\s+force|"
     r"broke\s+(?:open|the\s+lock)|forced\s+open"),

    ("entry_method", "shutter_lock_defeated",
     r"shutter(?:'s)?\s+lock|rolling\s+(?:gate|shutter)|shutter\s+was\s+cut"),
    ("entry_method", "steering_lock_broken",
     r"steering\s+lock|handle\s*lock"),

    ("target_type", "jewellery_retail",
     r"jewell?ery|jewell?er|gold\s+(?:showroom|shop|retail|outlet)|bullion"),
    ("target_type", "residential_street",
     r"parked\s+outside\s+his|residen(?:ce|tial)|outside\s+the\s+house|street"),

    ("premises_level", "ground_floor", r"ground[\s-]*floor|street[\s-]*level"),
    ("premises_level", "open_air",
     r"parked|open\s+(?:air|ground)|road\s*side|footpath"),

    ("cctv_defeated", "yes",
     r"(?:cctv|dvr|camera|surveillance)[^.\n]{0,60}?"
     r"(?:disconnect|disabl|damag|switched\s+off|tamper|inoperative|removed|cut)"
     r"|(?:disconnect|disabl|damag|tamper)[^.\n]{0,30}?(?:cctv|dvr|camera|surveillance)"),
    ("cctv_defeated", "no",
     r"(?:cctv|camera)[^.\n]{0,40}?(?:intact|working|recorded|functional)"),

    ("vault_breached", "no",
     r"(?:strong\s*room|vault|safe)[^.\n]{0,40}?"
     r"(?:not\s+(?:touched|breached|opened)|untouched|intact|secure|unopened)"),
    ("vault_breached", "yes",
     r"(?:strong\s*room|vault|safe)[^.\n]{0,40}?(?:broken|breached|opened|cut)"),

    ("goods_class", "gold_ornament",
     r"gold\s+(?:ornament|jewell?ery|article|chain|biscuit)|ornaments?\s+weighing"),
    ("goods_class", "vehicle",
     r"(?:motor\s*cycle|motorcycle|two[\s-]*wheeler|scooter|car)\s+"
     r"(?:was\s+)?(?:stolen|lifted|taken)"),

    ("escape_vehicle", "two_wheeler_no_plate",
     r"(?:motor\s*cycle|motorcycle|two[\s-]*wheeler|scooter)[^.\n]{0,60}?"
     r"(?:no\s+(?:number\s+)?plate|without\s+(?:any\s+)?(?:number\s+)?plate|"
     r"unregistered|no\s+registration)"),
    ("escape_vehicle", "stolen_vehicle_itself",
     r"(?:fled|escaped|left)\s+(?:on|with)\s+the\s+stolen"),
]

CREW_WORDS = {"one": "1", "two": "2", "three": "3", "four": "4", "five": "5"}
CREW_RE = re.compile(
    r"\b(one|two|three|four|five|\d)\s+(?:un(?:known|identified)\s+)?"
    r"(?:persons?|men|miscreants|accused|individuals|culprits|assailants)\b",
    re.IGNORECASE)

# "between 0215 and 0350 hours", "at about 2.30 am", "after midnight"
TIME_RE = re.compile(
    r"\b(?:(\d{3,4})\s*(?:hrs?|hours)|(\d{1,2})[.:](\d{2})\s*(a\.?m\.?|p\.?m\.?))",
    re.IGNORECASE)
MIDNIGHT_RE = re.compile(r"after\s+mid\s*night|in\s+the\s+dead\s+of\s+night",
                         re.IGNORECASE)


def _hour_of(text: str) -> tuple[int, tuple[int, int]] | None:
    m = TIME_RE.search(text)
    if m:
        if m.group(1):
            return int(m.group(1)[:-2] or 0), m.span()
        hour = int(m.group(2)) % 12
        if m.group(4).lower().startswith("p"):
            hour += 12
        return hour, m.span()
    m = MIDNIGHT_RE.search(text)
    if m:
        return 1, m.span()
    return None


PACK_LEXICON: list[tuple[int, str]] = [
    (10, r"kidnap|abduct|ransom|\b36[345]\b"),
    (6, r"smuggl|customs|consignment|\bDRI\b|bill\s+of\s+entry|undeclared|ICEGATE"),
    (2, r"launder|\bPMLA\b|hawala|layering|structuring|\bFIU\b"),
    (3, r"cyber|\bUPI\b|\bOTP\b|phishing|\b66\s*[CD]\b|\bNCRP\b|online\s+fraud"),
    (7, r"theft|burglar|robber|dacoit|house[\s-]*break|\b3(?:79|80|92|95)\b|\b457\b"),
]


# --------------------------------------------------------------------------
# The parse
# --------------------------------------------------------------------------

def parse_fir(text: str) -> dict[str, Any]:
    """Pull an intake submission out of an FIR document.

    Nothing here is committed. The result is a proposal for a human to check,
    and every value in it carries the text it was read from.
    """
    fields: dict[str, Any] = {}
    identifiers: dict[str, Any] = {}
    person: dict[str, Any] = {}
    mo: dict[str, Any] = {}

    # --- layer 1: the form
    district = _labelled(text, "district", r"distt\.?", r"dist\.?")
    if district:
        fields["district"] = _found(district[0][:60], district[1], text,
                                    "labelled field")

    officer = _labelled(text, r"investigating\s+officer", r"i\.?o\.?",
                        r"officer\s+in\s+charge", r"s\.?h\.?o\.?")
    if officer:
        fields["officer"] = _found(officer[0][:60], officer[1], text,
                                   "labelled field")

    narrative = _narrative(text)
    if narrative:
        body = re.sub(r"\s+", " ", narrative[0])[:4000]
        fields["narrative"] = _found(body, narrative[1], text,
                                     "narrative section")

    # A title is rarely a field on the form, so it is built from what is: the
    # offence and the place. An officer can overwrite it in the review.
    offence = _labelled(text, "offence", "offense", r"nature\s+of\s+(?:the\s+)?offence",
                        r"type\s+of\s+(?:the\s+)?case", "subject")
    if offence:
        fields["title"] = _found(offence[0][:120], offence[1], text,
                                 "labelled field")
    elif narrative:
        first = re.split(r"(?<=[.;])\s", narrative[0].strip())[0]
        fields["title"] = _found(re.sub(r"\s+", " ", first)[:110], narrative[1],
                                 text, "first sentence of the narrative")

    for label in ("date of occurrence", r"date\s+of\s+(?:the\s+)?(?:occurrence|offence|incident)",
                  r"occurrence\s+date", r"date\s+and\s+time\s+of\s+occurrence",
                  "date of report", r"fir\s+date", "date"):
        hit = _labelled(text, label)
        if hit:
            iso = _normalise_date(hit[0])
            if iso:
                fields["opened"] = _found(iso, hit[1], text,
                                          "date on the form")
                break

    name = _labelled(text, r"(?:name\s+of\s+(?:the\s+)?)?accused",
                     r"suspect(?:'s)?\s+name", r"name\s+of\s+(?:the\s+)?suspect")
    if name:
        # "Ramesh Yadav, s/o Bhanwar Lal, aged 39" -> "Ramesh Yadav"
        clean = re.split(r"\s*(?:,|\bs\s*/\s*o\b|\bson\s+of\b|\bd\s*/\s*o\b|\baged?\b|\(|\br\s*/\s*o\b)",
                         name[0], maxsplit=1)[0]
        clean = re.sub(r"\s+", " ", clean).strip(" .-")
        if 2 <= len(clean) <= 80:
            person["name"] = _found(clean, name[1], text, "labelled field")

    dob = _labelled(text, r"date\s+of\s+birth", r"d\.?o\.?b\.?")
    if dob:
        iso = _normalise_date(dob[0])
        if iso:
            person["dob"] = _found(iso, dob[1], text, "labelled field")

    # --- layer 2: identifier shapes
    masked = text
    m = IMEI_RE.search(text)
    if m:
        identifiers["imei"] = _found(m.group(1), m.span(1), text,
                                     "fifteen consecutive digits")
        masked = masked[: m.start(1)] + "#" * 15 + masked[m.end(1):]

    m = PHONE_RE.search(masked)
    if m:
        digits = re.sub(r"\D", "", m.group(1))
        identifiers["number"] = _found(f"+91 {digits[:5]} {digits[5:]}",
                                       m.span(1), text,
                                       "ten-digit mobile number")

    m = ACCOUNT_RE.search(text)
    if m:
        identifiers["account_no"] = _found(m.group(1), m.span(1), text,
                                           "digits following an account label")
        bank = _labelled(text, "bank", r"bank\s+name", r"name\s+of\s+(?:the\s+)?bank")
        if bank:
            identifiers["bank"] = _found(bank[0][:60], bank[1], text,
                                         "labelled field")

    m = REG_RE.search(text)
    if m:
        identifiers["registration"] = _found(
            re.sub(r"[\s-]+", " ", m.group(1)).upper(), m.span(1), text,
            "vehicle registration format")

    # --- layer 3: the lexicon
    body_text = fields.get("narrative", {}).get("value") or text
    for feature, value, pattern in MO_LEXICON:
        if feature in mo:
            continue
        m = re.search(pattern, body_text, re.IGNORECASE)
        if m:
            span = None
            idx = text.lower().find(m.group(0).lower())
            if idx >= 0:
                span = (idx, idx + len(m.group(0)))
            mo[feature] = _found(value, span, text,
                                 f"phrase “{m.group(0).strip()}”")

    m = CREW_RE.search(body_text)
    if m:
        word = m.group(1).lower()
        span = None
        idx = text.lower().find(m.group(0).lower())
        if idx >= 0:
            span = (idx, idx + len(m.group(0)))
        mo["crew_size"] = _found(CREW_WORDS.get(word, word), span, text,
                                 f"phrase “{m.group(0).strip()}”")

    hour = _hour_of(body_text)
    if hour:
        h, span = hour
        idx = None
        if span:
            fragment = body_text[span[0]:span[1]]
            found = text.lower().find(fragment.lower())
            idx = (found, found + len(fragment)) if found >= 0 else None
        band = "0200_0400" if 2 <= h < 4 else ("0000_0500" if h < 5 else None)
        if band:
            mo["time_band"] = _found(band, idx, text,
                                     f"time of occurrence around {h:02d}00")

    # --- crime pack
    for pack, pattern in PACK_LEXICON:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            fields["pack"] = _found(pack, m.span(), text,
                                    f"offence keyword “{m.group(0).strip()}”")
            break

    missing = [k for k in ("title", "district", "pack", "opened", "narrative")
               if k not in fields]
    return {
        "fields": fields,
        "person": person,
        "identifiers": identifiers,
        "mo_features": mo,
        "missing_required": missing,
        "counts": {
            "fields": len(fields), "identifiers": len(identifiers),
            "mo_features": len(mo), "person": len(person),
        },
    }
